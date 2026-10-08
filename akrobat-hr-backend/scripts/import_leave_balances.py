"""
Import the 2026 Singapore leave sheet into leave_balances.

    python scripts/import_leave_balances.py Leave_Record_for_App.xlsx            # dry run
    python scripts/import_leave_balances.py Leave_Record_for_App.xlsx --commit   # write

Run sql/034_leave_half_days.sql first (half-day values need numeric columns).

Mapping (sheet column -> leave_balances row for --year, default 2026):

  Annual Leave   total = Balance Forward 2025 + Annual Leave Entitlement
                 used  = Taken Annual Leave
                 remaining = total - used   (matches the sheet's own formula)
  MC             -> leave type "SICK LEAVE" (the app's name for medical leave)
                 total = MC Entitlement, used = MC (Used)
  Hospitalisation total = Hospitalization Entitlement, used = Hospitalization taken
  Childcare      only where the sheet has an entitlement (2 employees); also sets
                 employee_leave_tier to the 6 DAYS / 2 DAYS tier
  Maternity      not imported (sheet column is empty; HR/Super Admin manage it)
  Replacement    NOT imported unless --replacement-ph-date and --replacement-expiry
                 are given. The app stores one credit row per whole day, so
                 half-day figures (0.5) are reported and skipped.

Employees are matched to employees.full_name. Anything that doesn't match
exactly one employee is reported and skipped -- nothing is guessed.
"""

import argparse
import re
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

FIRST_DATA_ROW = 4
# column letters in the sheet
COL = dict(
    name="B",
    join="C",
    bf="D",
    ent="E",
    taken="G",
    sheet_bal="H",
    cc_ent="I",
    cc_taken="J",
    mc_ent="K",
    mc_used="L",
    hosp_ent="M",
    hosp_taken="N",
    repl="Q",
)


def num(v):
    """Cell -> float, or None if blank / non-numeric text."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"^\s*(\d+(?:\.\d+)?)", str(v))
    return float(m.group(1)) if m else None


def norm(s):
    return re.sub(r"[^A-Z0-9 ]", " ", (s or "").upper()).split()


def read_sheet(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        name = ws[f"{COL['name']}{r}"].value
        sn = ws[f"A{r}"].value
        # real employee rows have a numeric S/N; this skips blanks and the footnote
        if not name or not str(name).strip() or not isinstance(sn, (int, float)):
            continue
        g = lambda k: ws[f"{COL[k]}{r}"].value
        rows.append(
            dict(
                row=r,
                name=str(name).strip(),
                join=g("join"),
                bf=num(g("bf")) or 0.0,
                ent=num(g("ent")) or 0.0,
                taken=num(g("taken")) or 0.0,
                sheet_bal=num(g("sheet_bal")),
                cc_raw=g("cc_ent"),
                cc_ent=num(g("cc_ent")),
                cc_taken=num(g("cc_taken")) or 0.0,
                mc_ent=num(g("mc_ent")),
                mc_used=num(g("mc_used")) or 0.0,
                hosp_ent=num(g("hosp_ent")),
                hosp_taken=num(g("hosp_taken")) or 0.0,
                repl=num(g("repl")) or 0.0,
            )
        )
    return rows


def match_employee(name, employees):
    """Exactly-one match on full_name, else None + reason."""
    want = norm(name)
    # nickname after " - " is optional in the DB, so also try without it
    base = norm(name.split(" - ")[0])
    for tokens in (want, base):
        hits = [e for e in employees if norm(e["full_name"]) == tokens]
        if len(hits) == 1:
            return hits[0], None
        if len(hits) > 1:
            return None, "multiple employees with that name"
    # token-set fallback (word order / nickname differences)
    hits = [e for e in employees if set(base) <= set(norm(e["full_name"]))]
    if len(hits) == 1:
        return hits[0], None
    return None, "no match" if not hits else "ambiguous"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument(
        "--commit", action="store_true", help="write to the DB (default is a dry run)"
    )
    ap.add_argument(
        "--replacement-ph-date",
        help="public_holiday_date for imported replacement credits (YYYY-MM-DD)",
    )
    ap.add_argument(
        "--replacement-expiry",
        help="expiry_date for imported replacement credits (YYYY-MM-DD)",
    )
    args = ap.parse_args()

    from app.core.database import supabase_admin as db  # noqa: E402

    rows = read_sheet(args.xlsx)
    employees = (
        db.table("employees").select("id, full_name, joining_date").execute().data or []
    )
    types = {
        t["leave_name"]: t["id"]
        for t in db.table("leave_types").select("id, leave_name").execute().data
    }
    tiers = {
        t["tier_name"]: t["id"]
        for t in db.table("leave_policy_tiers")
        .select("id, tier_name, leave_type_id")
        .eq("leave_type_id", types.get("CHILDCARE LEAVE", ""))
        .execute()
        .data
        or []
    }

    for needed in (
        "ANNUAL LEAVE",
        "SICK LEAVE",
        "HOSPITALISATION LEAVE",
        "CHILDCARE LEAVE",
    ):
        if needed not in types:
            sys.exit(
                f"leave_types is missing {needed!r} -- run the 026 migration first."
            )

    plan, problems, skipped_repl = [], [], []
    for r in rows:
        emp, why = match_employee(r["name"], employees)
        if not emp:
            problems.append(f"row {r['row']}: {r['name']!r} -> {why}")
            continue

        # sanity: sheet's own balance formula
        calc = r["bf"] + r["ent"] - r["taken"]
        if r["sheet_bal"] is not None and abs(calc - r["sheet_bal"]) > 0.001:
            problems.append(
                f"row {r['row']}: {r['name']}: sheet balance {r['sheet_bal']} != BF+Ent-Taken {calc}"
            )

        # sanity: join date
        try:
            sheet_join = (
                datetime.strptime(str(r["join"]).strip(), "%d.%m.%Y").date()
                if r["join"]
                else None
            )
        except ValueError:
            sheet_join = None
        if (
            sheet_join
            and emp.get("joining_date")
            and str(emp["joining_date"])[:10] != sheet_join.isoformat()
        ):
            problems.append(
                f"row {r['row']}: {r['name']}: join date sheet {sheet_join} vs DB {emp['joining_date']}"
            )

        bal = lambda leave, total, used: dict(
            employee_id=emp["id"],
            leave_type_id=types[leave],
            year=args.year,
            total_days=total,
            used_days=used,
            remaining_days=total - used,
            _label=leave,
            _name=r["name"],
        )

        plan.append(bal("ANNUAL LEAVE", r["bf"] + r["ent"], r["taken"]))
        if r["mc_ent"] is not None:
            plan.append(bal("SICK LEAVE", r["mc_ent"], r["mc_used"]))
        if r["hosp_ent"] is not None:
            plan.append(bal("HOSPITALISATION LEAVE", r["hosp_ent"], r["hosp_taken"]))
        if r["cc_ent"] is not None:
            plan.append(bal("CHILDCARE LEAVE", r["cc_ent"], r["cc_taken"]))
            plan[-1]["_tier"] = f"{int(r['cc_ent'])} DAYS"
        if r["repl"]:
            if r["repl"] != int(r["repl"]):
                skipped_repl.append(
                    f"{r['name']}: {r['repl']} day(s) -- half days can't be stored as credits"
                )
            else:
                plan.append(
                    dict(
                        _replacement=int(r["repl"]),
                        employee_id=emp["id"],
                        _name=r["name"],
                    )
                )

    print(
        f"{len(rows)} employees in sheet, {len({p['employee_id'] for p in plan})} matched\n"
    )
    print(f"{'Employee':36} {'Leave':22} {'Total':>6} {'Used':>6} {'Left':>6}")
    for p in plan:
        if "_replacement" in p:
            print(
                f"{p['_name'][:35]:36} {'REPLACEMENT (credits)':22} {p['_replacement']:>6}"
            )
        else:
            print(
                f"{p['_name'][:35]:36} {p['_label']:22} {p['total_days']:>6g} {p['used_days']:>6g} {p['remaining_days']:>6g}"
            )
    for label, items in (
        ("NEEDS ATTENTION", problems),
        ("REPLACEMENT SKIPPED", skipped_repl),
    ):
        if items:
            print(f"\n{label}:")
            for i in items:
                print("  -", i)

    if not args.commit:
        print("\nDry run only. Re-run with --commit to write.")
        return
    if any("no match" in p or "ambiguous" in p or "multiple" in p for p in problems):
        sys.exit(
            "\nRefusing to commit while employees are unmatched. Fix names or the DB first."
        )

    repl_ok = args.replacement_ph_date and args.replacement_expiry
    for p in plan:
        if "_replacement" in p:
            if not repl_ok:
                continue
            for _ in range(p["_replacement"]):
                db.table("leave_replacement_credits").insert(
                    dict(
                        employee_id=p["employee_id"],
                        public_holiday_date=args.replacement_ph_date,
                        expiry_date=args.replacement_expiry,
                    )
                ).execute()
            continue
        row = {k: v for k, v in p.items() if not k.startswith("_")}
        db.table("leave_balances").upsert(
            row, on_conflict="employee_id,leave_type_id,year"
        ).execute()
        if "_tier" in p and p["_tier"] in tiers:
            db.table("employee_leave_tier").upsert(
                dict(
                    employee_id=p["employee_id"],
                    leave_type_id=p["leave_type_id"],
                    tier_id=tiers[p["_tier"]],
                ),
                on_conflict="employee_id,leave_type_id",
            ).execute()
    print("\nDone.")
    if not repl_ok:
        print(
            "Replacement credits were NOT written (pass --replacement-ph-date and --replacement-expiry)."
        )


if __name__ == "__main__":
    main()
